"""Read-only source validation helpers for the Day 3 Branch master."""

from __future__ import annotations

from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

EXPECTED_BRANCH_CODES = frozenset(
    {
        "BIL01",
        "BIN01",
        "BIP01",
        "BIR01",
        "BIS01",
        "BIY01",
        "BKL01",
        "BMU01",
        "BOL01",
        "SBO01",
    }
)


class BranchDataValidationError(ValueError):
    """Raised when a source file cannot safely support the Branch master."""


@dataclass(frozen=True)
class BranchSourceRecord:
    branch_code: str
    branch_name: str


@dataclass(frozen=True)
class BranchSourceValidation:
    records: tuple[BranchSourceRecord, ...]
    null_codes: int
    duplicate_codes: dict[str, int]
    missing_expected_codes: list[str]
    unexpected_codes: list[str]
    blank_names: int

    @property
    def is_valid(self) -> bool:
        return not any(
            (
                self.null_codes,
                self.duplicate_codes,
                self.missing_expected_codes,
                self.unexpected_codes,
                self.blank_names,
            )
        )


@dataclass(frozen=True)
class CoverageReport:
    distinct_source_codes: int
    matched_codes: int
    unmatched_codes: dict[str, int]

    @property
    def rows_affected_by_unmatched_codes(self) -> int:
        return sum(self.unmatched_codes.values())

    @property
    def is_valid(self) -> bool:
        return not self.unmatched_codes


def normalize_branch_code(value: object) -> str:
    """Apply the single approved branch-code normalization rule."""
    if value is None:
        return ""
    return str(value).strip().upper()


def _worksheet_rows(path: Path, sheet_name: str | None = None) -> tuple[list[str], list[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise BranchDataValidationError(
                    f"{path}: required worksheet {sheet_name!r} was not found."
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active

        rows = worksheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration as error:
            raise BranchDataValidationError(f"{path}: worksheet is empty.") from error

        headers = [str(value).strip() if value is not None else "" for value in header_values]
        return headers, list(rows)
    finally:
        workbook.close()


def _column_index(headers: list[str], column_name: str, path: Path) -> int:
    try:
        return headers.index(column_name)
    except ValueError as error:
        raise BranchDataValidationError(
            f"{path}: required column {column_name!r} was not found; available columns are {headers}."
        ) from error


def inspect_branch_profile(path: Path) -> BranchSourceValidation:
    """Read and validate branch rows without changing the source file or database."""
    headers, rows = _worksheet_rows(path, sheet_name="Branch Profile")
    code_index = _column_index(headers, "SHOP_CODE", path)
    name_index = _column_index(headers, "SHOP_NAME", path)

    records: list[BranchSourceRecord] = []
    null_codes = 0
    blank_names = 0
    code_counts: Counter[str] = Counter()

    for row in rows:
        code = normalize_branch_code(row[code_index] if code_index < len(row) else None)
        name_value = row[name_index] if name_index < len(row) else None
        name = str(name_value).strip() if name_value is not None else ""
        if not code:
            null_codes += 1
            continue
        if not name:
            blank_names += 1
        code_counts[code] += 1
        records.append(BranchSourceRecord(branch_code=code, branch_name=name))

    source_codes = set(code_counts)
    return BranchSourceValidation(
        records=tuple(records),
        null_codes=null_codes,
        duplicate_codes={code: count for code, count in sorted(code_counts.items()) if count > 1},
        missing_expected_codes=sorted(EXPECTED_BRANCH_CODES - source_codes),
        unexpected_codes=sorted(source_codes - EXPECTED_BRANCH_CODES),
        blank_names=blank_names,
    )


def coverage_report(path: Path, code_column: str, branch_codes: Iterable[str]) -> CoverageReport:
    """Report raw source codes that have no corresponding official Branch code."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration as error:
            raise BranchDataValidationError(f"{path}: worksheet is empty.") from error

        headers = [str(value).strip() if value is not None else "" for value in header_values]
        code_index = _column_index(headers, code_column, path)
        counts: Counter[str] = Counter()
        for row in rows:
            code = normalize_branch_code(row[code_index] if code_index < len(row) else None)
            counts[code] += 1
    finally:
        workbook.close()

    master_codes = {normalize_branch_code(code) for code in branch_codes}
    unmatched = {code: count for code, count in sorted(counts.items()) if code not in master_codes}
    return CoverageReport(
        distinct_source_codes=len(counts),
        matched_codes=len(set(counts) & master_codes),
        unmatched_codes=unmatched,
    )
