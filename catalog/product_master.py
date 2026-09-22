"""Source validation helpers for the Day 3 Product master."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

PRODUCT_SHEET_NAME = "product_data.csv"
REQUIRED_PRODUCT_COLUMNS = (
    "PRODUCT_CODE",
    "PRODUCT_DESC",
    "PRODUCT_FAMILY_ID",
    "PRODUCT_FAMILY_DESC",
    "PRICE",
)


class ProductDataValidationError(ValueError):
    """Raised when source data cannot safely form a Product master."""


@dataclass(frozen=True)
class ParsedPrice:
    value: Decimal | None
    valid: bool
    category: str


@dataclass(frozen=True)
class ProductSourceRecord:
    product_code: str
    canonical_product_id: str
    product_desc: str
    family_id: str
    family_desc: str
    list_price: Decimal | None
    price_valid: bool
    price_category: str

    def comparison_fields(self) -> tuple[object, ...]:
        """Fields that must agree if a source product code appears twice."""
        return (
            self.product_desc,
            self.family_id,
            self.family_desc,
            self.list_price,
            self.price_valid,
        )


@dataclass(frozen=True)
class ProductSourceValidation:
    records: tuple[ProductSourceRecord, ...]
    row_count: int
    missing_product_codes: int
    duplicate_codes: dict[str, int]
    conflicting_duplicate_codes: list[str]
    null_product_descriptions: int
    null_family_ids: int
    null_family_descriptions: int
    price_counts: dict[str, int]

    @property
    def is_loadable(self) -> bool:
        return not self.missing_product_codes and not self.conflicting_duplicate_codes

    @property
    def unique_records(self) -> tuple[ProductSourceRecord, ...]:
        """One record per code; duplicate records are known to be field-identical."""
        by_code: dict[str, ProductSourceRecord] = {}
        for record in self.records:
            by_code.setdefault(record.product_code, record)
        return tuple(by_code.values())


def normalize_product_code(value: object) -> str:
    """Apply the approved source-product identifier normalization rule."""
    if value is None:
        return ""
    return str(value).strip().upper()


def canonical_product_id_for(product_code: str) -> str:
    """Build the deterministic canonical identity for one normalized product code."""
    normalized_code = normalize_product_code(product_code)
    if not normalized_code:
        raise ProductDataValidationError(
            "A canonical product ID requires a product code."
        )
    return f"CP_{normalized_code}"


def parse_price(value: object) -> ParsedPrice:
    """Preserve real numeric prices while marking missing or invalid values unusable."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return ParsedPrice(value=None, valid=False, category="missing")

    try:
        parsed = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return ParsedPrice(value=None, valid=False, category="invalid")

    if not parsed.is_finite():
        return ParsedPrice(value=None, valid=False, category="invalid")
    if parsed > 0:
        return ParsedPrice(value=parsed, valid=True, category="positive")
    if parsed == 0:
        return ParsedPrice(value=parsed, valid=False, category="zero")
    return ParsedPrice(value=parsed, valid=False, category="negative")


def _cell_value(row: tuple[object, ...], index: int) -> object:
    return row[index] if index < len(row) else None


def _text_value(value: object) -> str:
    return str(value).strip() if value is not None else ""


def inspect_product_source(path: Path) -> ProductSourceValidation:
    """Read the product master workbook without changing its contents."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if PRODUCT_SHEET_NAME not in workbook.sheetnames:
            raise ProductDataValidationError(
                f"{path}: required worksheet {PRODUCT_SHEET_NAME!r} was not found."
            )
        worksheet = workbook[PRODUCT_SHEET_NAME]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as error:
            raise ProductDataValidationError(f"{path}: worksheet is empty.") from error

        headers = [_text_value(value) for value in header_row]
        missing_columns = sorted(set(REQUIRED_PRODUCT_COLUMNS) - set(headers))
        if missing_columns:
            raise ProductDataValidationError(
                f"{path}: required columns are missing: {missing_columns}."
            )
        indexes = {column: headers.index(column) for column in REQUIRED_PRODUCT_COLUMNS}

        records: list[ProductSourceRecord] = []
        missing_product_codes = 0
        null_product_descriptions = 0
        null_family_ids = 0
        null_family_descriptions = 0
        price_counts: Counter[str] = Counter()

        for row in rows:
            product_code = normalize_product_code(
                _cell_value(row, indexes["PRODUCT_CODE"])
            )
            if not product_code:
                missing_product_codes += 1
                continue

            product_desc = _text_value(_cell_value(row, indexes["PRODUCT_DESC"]))
            family_id = _text_value(_cell_value(row, indexes["PRODUCT_FAMILY_ID"]))
            family_desc = _text_value(_cell_value(row, indexes["PRODUCT_FAMILY_DESC"]))
            if not product_desc:
                null_product_descriptions += 1
            if not family_id:
                null_family_ids += 1
            if not family_desc:
                null_family_descriptions += 1

            parsed_price = parse_price(_cell_value(row, indexes["PRICE"]))
            price_counts[parsed_price.category] += 1
            records.append(
                ProductSourceRecord(
                    product_code=product_code,
                    canonical_product_id=canonical_product_id_for(product_code),
                    product_desc=product_desc,
                    family_id=family_id,
                    family_desc=family_desc,
                    list_price=parsed_price.value,
                    price_valid=parsed_price.valid,
                    price_category=parsed_price.category,
                )
            )
    finally:
        workbook.close()

    grouped_records: dict[str, list[ProductSourceRecord]] = {}
    for record in records:
        grouped_records.setdefault(record.product_code, []).append(record)
    duplicate_codes = {
        code: len(group)
        for code, group in sorted(grouped_records.items())
        if len(group) > 1
    }
    conflicting_duplicate_codes = [
        code
        for code, group in grouped_records.items()
        if len({record.comparison_fields() for record in group}) > 1
    ]

    return ProductSourceValidation(
        records=tuple(records),
        row_count=len(records) + missing_product_codes,
        missing_product_codes=missing_product_codes,
        duplicate_codes=duplicate_codes,
        conflicting_duplicate_codes=sorted(conflicting_duplicate_codes),
        null_product_descriptions=null_product_descriptions,
        null_family_ids=null_family_ids,
        null_family_descriptions=null_family_descriptions,
        price_counts=dict(sorted(price_counts.items())),
    )
