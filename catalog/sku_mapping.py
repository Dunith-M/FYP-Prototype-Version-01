"""Exact-code mapping and unresolved-queue helpers for Day 3 Steps 4 and 6."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from catalog.models import Product, SkuMapping

SALES_SHEET_NAME = "data"
REQUIRED_SALES_COLUMNS = ("SALES_PART", "TOTAL_UNITS")


class SalesSkuDataValidationError(ValueError):
    """Raised when sales source data cannot safely support SKU mapping."""


@dataclass(frozen=True)
class SalesCodeMetrics:
    row_count: int
    signed_units: Decimal
    absolute_units: Decimal


@dataclass(frozen=True)
class SalesSourceMetrics:
    total_rows: int
    blank_source_codes: int
    codes: dict[str, SalesCodeMetrics]
    missing_unit_values: int
    invalid_unit_values: int

    @property
    def distinct_codes(self) -> int:
        return len(self.codes)


@dataclass(frozen=True)
class UnresolvedSkuQueueItem:
    mapping: SkuMapping
    affected_sales_rows: int
    affected_signed_units: Decimal
    affected_absolute_units: Decimal


@dataclass(frozen=True)
class UnresolvedSkuQueueReport:
    total_distinct_sales_codes: int
    exact_matched_codes: int
    unmatched_source_codes: int
    affected_sales_rows: int
    affected_signed_units: Decimal
    blank_source_codes: int
    unresolved_mapping_count: int
    missing_unresolved_source_codes: tuple[str, ...]
    extra_unresolved_mapping_codes: tuple[str, ...]
    items: tuple[UnresolvedSkuQueueItem, ...]


def normalize_source_product_code(value: object) -> str:
    """Normalize source identifiers with only trim and uppercase."""
    if value is None:
        return ""
    return str(value).strip().upper()


def _parse_signed_units(value: object) -> tuple[Decimal | None, str]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing"
    try:
        parsed = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None, "invalid"
    if not parsed.is_finite():
        return None, "invalid"
    return parsed, "valid"


def extract_sales_source_metrics(path: Path) -> SalesSourceMetrics:
    """Stream only SALES_PART and TOTAL_UNITS from the raw historical sales file."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SALES_SHEET_NAME not in workbook.sheetnames:
            raise SalesSkuDataValidationError(
                f"{path}: required worksheet {SALES_SHEET_NAME!r} was not found."
            )
        worksheet = workbook[SALES_SHEET_NAME]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as error:
            raise SalesSkuDataValidationError(f"{path}: worksheet is empty.") from error

        headers = [
            str(value).strip() if value is not None else "" for value in header_row
        ]
        missing_columns = sorted(set(REQUIRED_SALES_COLUMNS) - set(headers))
        if missing_columns:
            raise SalesSkuDataValidationError(
                f"{path}: required columns are missing: {missing_columns}."
            )
        code_index = headers.index("SALES_PART")
        units_index = headers.index("TOTAL_UNITS")

        total_rows = 0
        blank_source_codes = 0
        missing_unit_values = 0
        invalid_unit_values = 0
        row_counts: Counter[str] = Counter()
        signed_unit_totals: Counter[str] = Counter()
        absolute_unit_totals: Counter[str] = Counter()

        for row in rows:
            total_rows += 1
            raw_code = row[code_index] if code_index < len(row) else None
            code = normalize_source_product_code(raw_code)
            raw_units = row[units_index] if units_index < len(row) else None
            units, unit_status = _parse_signed_units(raw_units)
            if unit_status == "missing":
                missing_unit_values += 1
            elif unit_status == "invalid":
                invalid_unit_values += 1

            if not code:
                blank_source_codes += 1
                continue

            row_counts[code] += 1
            if units is not None:
                signed_unit_totals[code] += units
                absolute_unit_totals[code] += abs(units)
    finally:
        workbook.close()

    return SalesSourceMetrics(
        total_rows=total_rows,
        blank_source_codes=blank_source_codes,
        codes={
            code: SalesCodeMetrics(
                row_count=row_counts[code],
                signed_units=signed_unit_totals[code],
                absolute_units=absolute_unit_totals[code],
            )
            for code in sorted(row_counts)
        },
        missing_unit_values=missing_unit_values,
        invalid_unit_values=invalid_unit_values,
    )


def get_unresolved_sales_skus():
    """Return the persistent, human-review queue without modifying mappings."""
    return SkuMapping.objects.filter(
        source_system=SkuMapping.SourceSystem.SALES,
        confidence_class=SkuMapping.ConfidenceClass.UNRESOLVED,
        canonical_product__isnull=True,
    ).order_by("source_product_code")


def build_unresolved_sales_queue(sales_file: Path) -> UnresolvedSkuQueueReport:
    """Combine unresolved mapping records with raw-sales impact metrics, read-only."""
    metrics = extract_sales_source_metrics(sales_file)
    product_codes = set(Product.objects.values_list("product_code", flat=True))
    source_codes = set(metrics.codes)
    unmatched_source_codes = source_codes - product_codes
    queue_mappings = list(get_unresolved_sales_skus())
    queue_codes = {mapping.source_product_code for mapping in queue_mappings}

    items = []
    for mapping in queue_mappings:
        code_metrics = metrics.codes.get(mapping.source_product_code)
        if code_metrics is None:
            code_metrics = SalesCodeMetrics(
                row_count=0,
                signed_units=Decimal(0),
                absolute_units=Decimal(0),
            )
        items.append(
            UnresolvedSkuQueueItem(
                mapping=mapping,
                affected_sales_rows=code_metrics.row_count,
                affected_signed_units=code_metrics.signed_units,
                affected_absolute_units=code_metrics.absolute_units,
            )
        )

    items.sort(
        key=lambda item: (
            -item.affected_sales_rows,
            -item.affected_absolute_units,
            item.mapping.source_product_code,
        )
    )
    unresolved_source_codes = unmatched_source_codes & queue_codes
    return UnresolvedSkuQueueReport(
        total_distinct_sales_codes=metrics.distinct_codes,
        exact_matched_codes=len(source_codes & product_codes),
        unmatched_source_codes=len(unmatched_source_codes),
        affected_sales_rows=sum(
            metrics.codes[code].row_count for code in unresolved_source_codes
        ),
        affected_signed_units=sum(
            (metrics.codes[code].signed_units for code in unresolved_source_codes),
            start=Decimal(0),
        ),
        blank_source_codes=metrics.blank_source_codes,
        unresolved_mapping_count=len(queue_mappings),
        missing_unresolved_source_codes=tuple(
            sorted(unmatched_source_codes - queue_codes)
        ),
        extra_unresolved_mapping_codes=tuple(
            sorted(queue_codes - unmatched_source_codes)
        ),
        items=tuple(items),
    )
