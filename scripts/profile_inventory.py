from collections import defaultdict
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
INVENTORY_FILE = PROJECT_ROOT / "data" / "raw" / "inventory_data.xlsx"
SHEET_NAME = "inventory_data.csv"
OUTPUT_FILE = PROJECT_ROOT / "data" / "artifacts" / "inventory_profile_report.xlsx"


def format_sheet(worksheet):
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        width = min(
            max(len(str(cell.value or "")) for cell in column_cells) + 2,
            40,
        )
        worksheet.column_dimensions[column_cells[0].column_letter].width = width


def as_number(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def as_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return pd.Timestamp(value).date()
    converted = pd.to_datetime(value, errors="coerce")
    return None if pd.isna(converted) else converted.date()


def main():
    if not INVENTORY_FILE.exists():
        raise FileNotFoundError(f"Inventory file was not found:\n{INVENTORY_FILE}")

    # Stream the large workbook row by row to avoid using excessive memory.
    workbook = load_workbook(INVENTORY_FILE, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        workbook.close()
        raise ValueError(f"Sheet '{SHEET_NAME}' was not found. Available sheets: {available}")

    worksheet = workbook[SHEET_NAME]
    rows = worksheet.iter_rows(values_only=True)
    headers = list(next(rows))
    required_columns = {"AS_AT_DATE", "PRODUCT_CODE", "SITE_ID", "QTY"}
    missing_columns = required_columns - set(headers)
    if missing_columns:
        workbook.close()
        raise ValueError(
            "Missing required columns: " + ", ".join(sorted(missing_columns))
        )

    indexes = {name: headers.index(name) for name in required_columns}
    missing_counts = {header: 0 for header in headers}
    type_samples = {header: None for header in headers}
    unique_products = set()
    unique_sites = set()
    by_site = defaultdict(lambda: {"Records": 0, "Products": set(), "Total_QTY": 0})
    by_date = defaultdict(
        lambda: {"Records": 0, "Products": set(), "Sites": set(), "Total_QTY": 0}
    )
    seen_rows = set()
    sample_rows = []
    row_count = total_qty = 0
    negative_qty = zero_qty = positive_qty = duplicate_rows = 0

    for raw_values in rows:
        values = list(raw_values)
        if not any(value is not None for value in values):
            continue
        values.extend([None] * (len(headers) - len(values)))
        values = values[: len(headers)]
        row_count += 1

        if len(sample_rows) < 100:
            sample_rows.append(dict(zip(headers, values)))

        row_key = sha256(repr(values).encode("utf-8")).digest()
        if row_key in seen_rows:
            duplicate_rows += 1
        else:
            seen_rows.add(row_key)

        for header, value in zip(headers, values):
            if value is None or value == "":
                missing_counts[header] += 1
            elif type_samples[header] is None:
                type_samples[header] = type(value).__name__

        product = values[indexes["PRODUCT_CODE"]]
        site = values[indexes["SITE_ID"]]
        inventory_date = as_date(values[indexes["AS_AT_DATE"]])
        quantity = as_number(values[indexes["QTY"]])
        product = str(product) if product not in (None, "") else None
        site = str(site) if site not in (None, "") else None

        if product:
            unique_products.add(product)
        if site:
            unique_sites.add(site)

        site_key = site or "(blank)"
        by_site[site_key]["Records"] += 1
        if product:
            by_site[site_key]["Products"].add(product)

        date_key = str(inventory_date) if inventory_date else "(invalid/blank)"
        by_date[date_key]["Records"] += 1
        if product:
            by_date[date_key]["Products"].add(product)
        if site:
            by_date[date_key]["Sites"].add(site)

        if quantity is not None:
            total_qty += quantity
            by_site[site_key]["Total_QTY"] += quantity
            by_date[date_key]["Total_QTY"] += quantity
            if quantity < 0:
                negative_qty += 1
            elif quantity == 0:
                zero_qty += 1
            else:
                positive_qty += 1

    workbook.close()

    valid_dates = [key for key in by_date if key != "(invalid/blank)"]
    summary = pd.DataFrame(
        {
            "Metric": [
                "Source file", "Sheet", "Rows", "Columns", "Unique products",
                "Unique sites", "Minimum AS_AT_DATE", "Maximum AS_AT_DATE",
                "Total quantity", "Negative QTY rows", "Zero QTY rows",
                "Positive QTY rows", "Exact duplicate rows",
            ],
            "Value": [
                INVENTORY_FILE.name, SHEET_NAME, row_count, len(headers),
                len(unique_products), len(unique_sites), min(valid_dates, default=""),
                max(valid_dates, default=""), total_qty, negative_qty, zero_qty,
                positive_qty, duplicate_rows,
            ],
        }
    )

    data_types = pd.DataFrame(
        {"Column": headers, "Data type": [type_samples[h] or "empty" for h in headers]}
    )
    missing_values = pd.DataFrame(
        {"Column": headers, "Missing values": [missing_counts[h] for h in headers]}
    )
    checks = pd.DataFrame(
        {
            "Check": [
                "Negative QTY", "Zero QTY", "Positive QTY", "Exact duplicate rows",
                "Invalid or blank AS_AT_DATE",
            ],
            "Count": [
                negative_qty, zero_qty, positive_qty, duplicate_rows,
                by_date.get("(invalid/blank)", {}).get("Records", 0),
            ],
        }
    )
    by_site_report = pd.DataFrame(
        [
            {
                "SITE_ID": site,
                "Records": values["Records"],
                "Unique_Products": len(values["Products"]),
                "Total_QTY": values["Total_QTY"],
            }
            for site, values in by_site.items()
        ]
    ).sort_values("Total_QTY", ascending=False)
    by_date_report = pd.DataFrame(
        [
            {
                "AS_AT_DATE": inventory_date,
                "Records": values["Records"],
                "Unique_Products": len(values["Products"]),
                "Unique_Sites": len(values["Sites"]),
                "Total_QTY": values["Total_QTY"],
            }
            for inventory_date, values in by_date.items()
        ]
    ).sort_values("AS_AT_DATE")

    reports = {
        "Summary": summary,
        "Data_Types": data_types,
        "Missing_Values": missing_values,
        "Checks": checks,
        "By_Site": by_site_report,
        "By_Date": by_date_report,
        "Sample_Rows": pd.DataFrame(sample_rows, columns=headers),
    }

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        for sheet_name, report in reports.items():
            report.to_excel(writer, sheet_name=sheet_name, index=False)
            format_sheet(writer.sheets[sheet_name])

    print("=" * 80)
    print("INVENTORY DATA PROFILING COMPLETE")
    print("=" * 80)
    print(f"Source: {INVENTORY_FILE}")
    print(f"Rows: {row_count:,}")
    print(f"Columns: {len(headers)}")
    print(f"Unique products: {len(unique_products):,}")
    print(f"Unique sites: {len(unique_sites):,}")
    print(f"Total quantity: {total_qty:,.0f}")
    print(f"Downloadable report: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
