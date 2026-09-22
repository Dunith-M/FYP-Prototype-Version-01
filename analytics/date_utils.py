"""Shared, source-preserving date normalization for the historical MVP data."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


ASIA_COLOMBO = ZoneInfo("Asia/Colombo")


class DateSourceValidationError(ValueError):
    """Raised when a required workbook sheet or date column is unavailable."""


@dataclass(frozen=True)
class DateNormalizationResult:
    """The untouched source value and its safely derived date, if parseable."""

    original_value: Any
    normalized_date: date | None
    status: str


@dataclass(frozen=True)
class DateSourceProfile:
    source_name: str
    source_column: str
    total_rows: int
    parsed_dates: int
    null_dates: int
    unparseable_dates: int
    minimum_date: date | None
    maximum_date: date | None
    unique_date_count: int
    duplicate_date_rows: int
    date_row_counts: dict[date, int]

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_name": self.source_name,
            "source_column": self.source_column,
            "total_rows": self.total_rows,
            "parsed_dates": self.parsed_dates,
            "null_dates": self.null_dates,
            "unparseable_dates": self.unparseable_dates,
            "minimum_date": self.minimum_date.isoformat()
            if self.minimum_date
            else None,
            "maximum_date": self.maximum_date.isoformat()
            if self.maximum_date
            else None,
            "unique_date_count": self.unique_date_count,
            "duplicate_date_rows": self.duplicate_date_rows,
        }


@dataclass(frozen=True)
class DateNormalizationReport:
    sales: DateSourceProfile
    inventory: DateSourceProfile
    calendar: DateSourceProfile
    sales_dates_missing_from_calendar: tuple[date, ...]
    sales_rows_missing_from_calendar: int
    inventory_dates_missing_from_calendar: tuple[date, ...]
    inventory_rows_missing_from_calendar: int

    def to_dict(self) -> dict[str, Any]:
        return {
            "sales": self.sales.to_dict(),
            "inventory": self.inventory.to_dict(),
            "calendar": self.calendar.to_dict(),
            "calendar_coverage": {
                "sales_dates_missing_from_calendar": [
                    value.isoformat() for value in self.sales_dates_missing_from_calendar
                ],
                "sales_rows_missing_from_calendar": self.sales_rows_missing_from_calendar,
                "inventory_dates_missing_from_calendar": [
                    value.isoformat()
                    for value in self.inventory_dates_missing_from_calendar
                ],
                "inventory_rows_missing_from_calendar": self.inventory_rows_missing_from_calendar,
            },
        }


def _date_from_datetime(value: datetime) -> date:
    """Respect explicit offsets but never reinterpret naive source timestamps as UTC."""

    if value.tzinfo is not None and value.utcoffset() is not None:
        return value.astimezone(ASIA_COLOMBO).date()
    return value.date()


def _parse_string_date(value: str) -> date | None:
    """Parse ISO-style source values without guessing ambiguous day/month ordering."""

    normalized = value.strip()
    if not normalized:
        return None

    try:
        return date.fromisoformat(normalized)
    except ValueError:
        pass

    try:
        return _date_from_datetime(datetime.fromisoformat(normalized.replace("Z", "+00:00")))
    except ValueError:
        pass

    for pattern in ("%Y/%m/%d", "%Y/%m/%d %H:%M:%S"):
        try:
            return _date_from_datetime(datetime.strptime(normalized, pattern))
        except ValueError:
            continue
    return None


def normalize_source_date(value: Any) -> DateNormalizationResult:
    """Return a date without changing the supplied raw source value.

    A naive timestamp is treated as a Sri Lankan business-local source timestamp;
    it is not converted from UTC.  Offset-aware values are converted to
    ``Asia/Colombo`` before their date is derived.
    """

    if value is None:
        return DateNormalizationResult(value, None, "NULL")
    if isinstance(value, datetime):
        return DateNormalizationResult(value, _date_from_datetime(value), "PARSED")
    if isinstance(value, date):
        return DateNormalizationResult(value, value, "PARSED")
    if isinstance(value, str):
        parsed = _parse_string_date(value)
        if not value.strip():
            return DateNormalizationResult(value, None, "NULL")
        if parsed is not None:
            return DateNormalizationResult(value, parsed, "PARSED")
    return DateNormalizationResult(value, None, "UNPARSEABLE")


def normalize_sales_date(value: Any) -> DateNormalizationResult:
    """Derive ``business_date`` from the preserved ``SALES_DATE`` value."""

    return normalize_source_date(value)


def normalize_inventory_date(value: Any) -> DateNormalizationResult:
    """Derive ``snapshot_date`` from ``AS_AT_DATE``."""

    return normalize_source_date(value)


def normalize_calendar_date(value: Any) -> DateNormalizationResult:
    """Derive ``calendar_date`` from the approved calendar date field."""

    return normalize_source_date(value)


def year_month_from_calendar_date(calendar_date: date) -> str:
    """The sole YYYY-MM derivation for calendar-based future monthly logic."""

    return calendar_date.strftime("%Y-%m")


def profile_workbook_date_column(
    workbook_path: Path,
    sheet_name: str,
    source_name: str,
    column_name: str,
    normalizer: Callable[[Any], DateNormalizationResult],
) -> DateSourceProfile:
    """Stream one source date column and calculate non-mutating quality metrics."""

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise DateSourceValidationError(
                f"{source_name}: required sheet {sheet_name!r} was not found in {workbook_path}."
            )
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as error:
            raise DateSourceValidationError(
                f"{source_name}: sheet {sheet_name!r} has no header row."
            ) from error
        try:
            column_index = header.index(column_name)
        except ValueError as error:
            raise DateSourceValidationError(
                f"{source_name}: required date column {column_name!r} was not found."
            ) from error

        total_rows = parsed_dates = null_dates = unparseable_dates = 0
        date_row_counts: Counter[date] = Counter()
        for row in rows:
            total_rows += 1
            value = row[column_index] if column_index < len(row) else None
            result = normalizer(value)
            if result.status == "PARSED":
                parsed_dates += 1
                date_row_counts[result.normalized_date] += 1
            elif result.status == "NULL":
                null_dates += 1
            else:
                unparseable_dates += 1
    finally:
        workbook.close()

    observed_dates = tuple(sorted(date_row_counts))
    return DateSourceProfile(
        source_name=source_name,
        source_column=column_name,
        total_rows=total_rows,
        parsed_dates=parsed_dates,
        null_dates=null_dates,
        unparseable_dates=unparseable_dates,
        minimum_date=observed_dates[0] if observed_dates else None,
        maximum_date=observed_dates[-1] if observed_dates else None,
        unique_date_count=len(observed_dates),
        duplicate_date_rows=sum(count - 1 for count in date_row_counts.values()),
        date_row_counts=dict(date_row_counts),
    )


def normalize_date_sources(
    sales_workbook: Path,
    inventory_workbook: Path,
    calendar_workbook: Path,
) -> DateNormalizationReport:
    """Normalize/profiles dates only; it deliberately performs no joins or writes."""

    sales = profile_workbook_date_column(
        sales_workbook, "data", "sales", "SALES_DATE", normalize_sales_date
    )
    inventory = profile_workbook_date_column(
        inventory_workbook,
        "inventory_data.csv",
        "inventory",
        "AS_AT_DATE",
        normalize_inventory_date,
    )
    calendar = profile_workbook_date_column(
        calendar_workbook, "Sheet1", "calendar", "DATE", normalize_calendar_date
    )

    calendar_dates = set(calendar.date_row_counts)
    sales_missing = tuple(sorted(set(sales.date_row_counts) - calendar_dates))
    inventory_missing = tuple(sorted(set(inventory.date_row_counts) - calendar_dates))
    return DateNormalizationReport(
        sales=sales,
        inventory=inventory,
        calendar=calendar,
        sales_dates_missing_from_calendar=sales_missing,
        sales_rows_missing_from_calendar=sum(
            sales.date_row_counts[value] for value in sales_missing
        ),
        inventory_dates_missing_from_calendar=inventory_missing,
        inventory_rows_missing_from_calendar=sum(
            inventory.date_row_counts[value] for value in inventory_missing
        ),
    )
