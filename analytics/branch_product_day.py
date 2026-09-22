"""Create validated daily inventory and safely left-join it to SALES_DAY."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from analytics.date_utils import normalize_inventory_date
from analytics.sales_day import SALES_DAY_KEY
from catalog.branch_master import normalize_branch_code
from catalog.models import Branch, Product
from catalog.product_master import normalize_product_code


INVENTORY_SHEET_NAME = "inventory_data.csv"
REQUIRED_INVENTORY_COLUMNS = ("AS_AT_DATE", "PRODUCT_CODE", "SITE_ID", "QTY")
INVENTORY_DAY_KEY = ("branch_code", "canonical_product_id", "snapshot_date")
INVENTORY_DAY_SCHEMA = pa.schema(
    [
        pa.field("branch_code", pa.string()),
        pa.field("canonical_product_id", pa.string()),
        pa.field("snapshot_date", pa.date32()),
        pa.field("inventory_qty", pa.float64()),
        pa.field("availability_state", pa.string()),
    ]
)
UNRESOLVED_INVENTORY_SCHEMA = pa.schema(
    [
        pa.field("source_row_number", pa.int64()),
        pa.field("AS_AT_DATE", pa.string()),
        pa.field("SITE_ID", pa.string()),
        pa.field("PRODUCT_CODE", pa.string()),
        pa.field("QTY", pa.float64()),
        pa.field("branch_code", pa.string()),
        pa.field("source_product_code", pa.string()),
        pa.field("snapshot_date", pa.date32()),
        pa.field("unresolved_reason", pa.string()),
    ]
)


class InventoryDayValidationError(ValueError):
    """Raised when inventory cannot safely support the daily left join."""


@dataclass(frozen=True)
class InventoryDayMetrics:
    source_inventory_rows: int
    resolved_inventory_rows: int
    unresolved_product_rows: int
    unresolved_product_codes: tuple[str, ...]
    unmatched_branch_rows: int
    unmatched_branch_codes: dict[str, int]
    duplicate_inventory_keys: int


@dataclass(frozen=True)
class InventoryDayBuildResult:
    metrics: InventoryDayMetrics
    output_path: Path
    unresolved_output_path: Path


@dataclass(frozen=True)
class BranchProductDayMetrics:
    sales_day_rows: int
    branch_product_day_rows: int
    inventory_match_rows: int
    unknown_inventory_rows: int
    duplicate_sales_day_keys: int
    duplicate_inventory_day_keys: int
    raw_sales_units: float
    joined_sales_units: float
    raw_sales_value: float
    joined_sales_value: float

    @property
    def unit_difference(self) -> float:
        return self.raw_sales_units - self.joined_sales_units

    @property
    def value_difference(self) -> float:
        return self.raw_sales_value - self.joined_sales_value


@dataclass(frozen=True)
class BranchProductDayBuildResult:
    branch_product_day: pd.DataFrame
    metrics: BranchProductDayMetrics


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normalize_branch(value: object) -> str:
    return normalize_branch_code(value)


def _normalize_product(value: object) -> str:
    return normalize_product_code(value)


def _parse_inventory_qty(value: object) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise InventoryDayValidationError("Inventory QTY is missing.")
    try:
        quantity = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError) as error:
        raise InventoryDayValidationError(f"Inventory QTY is unparseable: {value!r}.") from error
    if quantity != quantity or quantity in (float("inf"), float("-inf")):
        raise InventoryDayValidationError(f"Inventory QTY is non-finite: {value!r}.")
    if quantity < 0:
        raise InventoryDayValidationError(
            f"Inventory QTY is negative and has no approved availability-state rule: {value!r}."
        )
    return quantity


def _availability_state(quantity: float) -> str:
    return "AVAILABLE" if quantity > 0 else "ZERO_CONFIRMED"


def _write_batch(writer: pq.ParquetWriter, schema: pa.Schema, columns: dict[str, list[Any]]) -> None:
    if columns[next(iter(columns))]:
        writer.write_table(pa.Table.from_pydict(columns, schema=schema))


def _empty_columns(schema: pa.Schema) -> dict[str, list[Any]]:
    return {field.name: [] for field in schema}


def _flush_resolved(
    writer: pq.ParquetWriter, columns: dict[str, list[Any]]
) -> None:
    _write_batch(writer, INVENTORY_DAY_SCHEMA, columns)
    for values in columns.values():
        values.clear()


def _flush_unresolved(
    writer: pq.ParquetWriter, columns: dict[str, list[Any]]
) -> None:
    _write_batch(writer, UNRESOLVED_INVENTORY_SCHEMA, columns)
    for values in columns.values():
        values.clear()


def _temporary_path(path: Path) -> Path:
    return path.with_suffix(path.suffix + ".tmp")


def _product_codes_to_canonical_ids() -> dict[str, str]:
    return {
        normalize_product_code(product_code): canonical_product_id
        for product_code, canonical_product_id in Product.objects.values_list(
            "product_code", "canonical_product_id"
        )
    }


def build_inventory_day(
    inventory_file: Path,
    output_path: Path,
    unresolved_output_path: Path,
    batch_size: int = 50_000,
) -> InventoryDayBuildResult:
    """Stream validated inventory snapshots into unique daily Parquet records.

    Duplicate snapshot keys are rejected rather than summed.  Product codes not
    present in the Product master are preserved separately for audit.
    """

    output_path.parent.mkdir(parents=True, exist_ok=True)
    unresolved_output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_output = _temporary_path(output_path)
    temporary_unresolved = _temporary_path(unresolved_output_path)
    temporary_output.unlink(missing_ok=True)
    temporary_unresolved.unlink(missing_ok=True)

    official_branches = set(Branch.objects.values_list("branch_code", flat=True))
    canonical_by_product_code = _product_codes_to_canonical_ids()
    seen_keys: set[tuple[str, str, object]] = set()
    duplicate_keys: Counter[tuple[str, str, object]] = Counter()
    unmatched_branches: Counter[str] = Counter()
    unresolved_codes: set[str] = set()
    source_rows = resolved_rows = unresolved_rows = 0
    resolved_columns = _empty_columns(INVENTORY_DAY_SCHEMA)
    unresolved_columns = _empty_columns(UNRESOLVED_INVENTORY_SCHEMA)

    try:
        workbook = load_workbook(inventory_file, read_only=True, data_only=True)
        try:
            if INVENTORY_SHEET_NAME not in workbook.sheetnames:
                raise InventoryDayValidationError(
                    f"{inventory_file}: required worksheet {INVENTORY_SHEET_NAME!r} was not found."
                )
            worksheet = workbook[INVENTORY_SHEET_NAME]
            rows = worksheet.iter_rows(values_only=True)
            try:
                header = tuple(_text(value) for value in next(rows))
            except StopIteration as error:
                raise InventoryDayValidationError(
                    f"{inventory_file}: inventory worksheet is empty."
                ) from error
            missing_columns = sorted(set(REQUIRED_INVENTORY_COLUMNS) - set(header))
            if missing_columns:
                raise InventoryDayValidationError(
                    f"{inventory_file}: required inventory columns are missing: {missing_columns}."
                )
            index = {column: header.index(column) for column in REQUIRED_INVENTORY_COLUMNS}

            with pq.ParquetWriter(temporary_output, INVENTORY_DAY_SCHEMA) as resolved_writer, pq.ParquetWriter(
                temporary_unresolved, UNRESOLVED_INVENTORY_SCHEMA
            ) as unresolved_writer:
                for source_row_number, row in enumerate(rows, start=2):
                    source_rows += 1
                    raw_branch = row[index["SITE_ID"]] if index["SITE_ID"] < len(row) else None
                    raw_product = row[index["PRODUCT_CODE"]] if index["PRODUCT_CODE"] < len(row) else None
                    raw_date = row[index["AS_AT_DATE"]] if index["AS_AT_DATE"] < len(row) else None
                    raw_qty = row[index["QTY"]] if index["QTY"] < len(row) else None
                    branch_code = _normalize_branch(raw_branch)
                    source_product_code = _normalize_product(raw_product)
                    parsed_date = normalize_inventory_date(raw_date)

                    if not branch_code:
                        unmatched_branches[""] += 1
                        continue
                    if branch_code not in official_branches:
                        unmatched_branches[branch_code] += 1
                        continue
                    if not source_product_code:
                        raise InventoryDayValidationError(
                            f"Inventory PRODUCT_CODE is blank at source row {source_row_number}."
                        )
                    if parsed_date.normalized_date is None:
                        raise InventoryDayValidationError(
                            "Inventory AS_AT_DATE is invalid at source row "
                            f"{source_row_number}; status={parsed_date.status}."
                        )
                    quantity = _parse_inventory_qty(raw_qty)
                    canonical_product_id = canonical_by_product_code.get(source_product_code)
                    if canonical_product_id is None:
                        unresolved_rows += 1
                        unresolved_codes.add(source_product_code)
                        unresolved_columns["source_row_number"].append(source_row_number)
                        unresolved_columns["AS_AT_DATE"].append(_text(raw_date))
                        unresolved_columns["SITE_ID"].append(_text(raw_branch))
                        unresolved_columns["PRODUCT_CODE"].append(_text(raw_product))
                        unresolved_columns["QTY"].append(quantity)
                        unresolved_columns["branch_code"].append(branch_code)
                        unresolved_columns["source_product_code"].append(source_product_code)
                        unresolved_columns["snapshot_date"].append(parsed_date.normalized_date)
                        unresolved_columns["unresolved_reason"].append(
                            "Inventory product code absent from Product master"
                        )
                        if len(unresolved_columns["source_row_number"]) >= batch_size:
                            _flush_unresolved(unresolved_writer, unresolved_columns)
                        continue

                    key = (branch_code, canonical_product_id, parsed_date.normalized_date)
                    if key in seen_keys:
                        duplicate_keys[key] += 1
                        continue
                    seen_keys.add(key)
                    resolved_rows += 1
                    resolved_columns["branch_code"].append(branch_code)
                    resolved_columns["canonical_product_id"].append(canonical_product_id)
                    resolved_columns["snapshot_date"].append(parsed_date.normalized_date)
                    resolved_columns["inventory_qty"].append(quantity)
                    resolved_columns["availability_state"].append(_availability_state(quantity))
                    if len(resolved_columns["branch_code"]) >= batch_size:
                        _flush_resolved(resolved_writer, resolved_columns)

                _flush_resolved(resolved_writer, resolved_columns)
                _flush_unresolved(unresolved_writer, unresolved_columns)
        finally:
            workbook.close()

        if unmatched_branches:
            raise InventoryDayValidationError(
                "Inventory branch mapping is blocked: SITE_ID values absent from "
                f"Branch master: {dict(sorted(unmatched_branches.items()))}."
            )
        if duplicate_keys:
            sample = [
                {"key": key, "extra_rows": count}
                for key, count in list(sorted(duplicate_keys.items()))[:10]
            ]
            raise InventoryDayValidationError(
                "Inventory snapshot grain is blocked: duplicate branch/product/date "
                f"keys={sum(duplicate_keys.values())}; sample={sample}."
            )

        temporary_output.replace(output_path)
        temporary_unresolved.replace(unresolved_output_path)
    except Exception:
        temporary_output.unlink(missing_ok=True)
        temporary_unresolved.unlink(missing_ok=True)
        raise

    return InventoryDayBuildResult(
        metrics=InventoryDayMetrics(
            source_inventory_rows=source_rows,
            resolved_inventory_rows=resolved_rows,
            unresolved_product_rows=unresolved_rows,
            unresolved_product_codes=tuple(sorted(unresolved_codes)),
            unmatched_branch_rows=sum(unmatched_branches.values()),
            unmatched_branch_codes=dict(sorted(unmatched_branches.items())),
            duplicate_inventory_keys=sum(duplicate_keys.values()),
        ),
        output_path=output_path,
        unresolved_output_path=unresolved_output_path,
    )


def _validate_sales_day(frame: pd.DataFrame) -> int:
    missing_columns = sorted(
        set((*SALES_DAY_KEY, "sales_units", "sales_value")) - set(frame.columns)
    )
    if missing_columns:
        raise InventoryDayValidationError(
            f"SALES_DAY is missing required columns: {missing_columns}."
        )
    if frame[list(SALES_DAY_KEY)].isna().any().any():
        raise InventoryDayValidationError("SALES_DAY has null grain values.")
    duplicates = int(frame.duplicated(list(SALES_DAY_KEY)).sum())
    if duplicates:
        raise InventoryDayValidationError(f"SALES_DAY duplicate keys={duplicates}.")
    return duplicates


def _validate_inventory_day(frame: pd.DataFrame) -> int:
    missing_columns = sorted(
        set((*INVENTORY_DAY_KEY, "inventory_qty", "availability_state"))
        - set(frame.columns)
    )
    if missing_columns:
        raise InventoryDayValidationError(
            f"INVENTORY_DAY is missing required columns: {missing_columns}."
        )
    if frame[list(INVENTORY_DAY_KEY)].isna().any().any():
        raise InventoryDayValidationError("INVENTORY_DAY has null grain values.")
    duplicates = int(frame.duplicated(list(INVENTORY_DAY_KEY)).sum())
    if duplicates:
        raise InventoryDayValidationError(f"INVENTORY_DAY duplicate keys={duplicates}.")
    return duplicates


def build_branch_product_day(
    sales_day_path: Path, inventory_day_path: Path
) -> BranchProductDayBuildResult:
    """Left-join validated snapshot observations to SALES_DAY at exact date equality."""

    sales_day = pd.read_parquet(sales_day_path)
    inventory_day = pd.read_parquet(inventory_day_path)
    sales_duplicates = _validate_sales_day(sales_day)
    inventory_duplicates = _validate_inventory_day(inventory_day)

    joined = sales_day.merge(
        inventory_day,
        how="left",
        left_on=list(SALES_DAY_KEY),
        right_on=["branch_code", "canonical_product_id", "snapshot_date"],
        validate="one_to_one",
        indicator=True,
    )
    if len(joined) != len(sales_day):
        raise InventoryDayValidationError(
            "Left join changed SALES_DAY row count: "
            f"before={len(sales_day)} after={len(joined)}."
        )
    matched = joined["_merge"] == "both"
    joined.loc[~matched, "availability_state"] = "UNKNOWN"
    if joined.loc[~matched, "inventory_qty"].notna().any():
        raise InventoryDayValidationError(
            "UNKNOWN inventory rows must retain null inventory_qty."
        )
    joined = joined.drop(columns=["snapshot_date", "_merge"])
    joined = joined.sort_values(list(SALES_DAY_KEY), kind="stable").reset_index(drop=True)

    metrics = BranchProductDayMetrics(
        sales_day_rows=len(sales_day),
        branch_product_day_rows=len(joined),
        inventory_match_rows=int(matched.sum()),
        unknown_inventory_rows=int((~matched).sum()),
        duplicate_sales_day_keys=sales_duplicates,
        duplicate_inventory_day_keys=inventory_duplicates,
        raw_sales_units=float(sales_day["sales_units"].sum()),
        joined_sales_units=float(joined["sales_units"].sum()),
        raw_sales_value=float(sales_day["sales_value"].sum()),
        joined_sales_value=float(joined["sales_value"].sum()),
    )
    return BranchProductDayBuildResult(joined, metrics)


def write_branch_product_day(result: BranchProductDayBuildResult, output_path: Path) -> None:
    """Overwrite the single deterministic left-join output for the current inputs."""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    result.branch_product_day.to_parquet(output_path, index=False)
